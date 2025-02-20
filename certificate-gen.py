import os
import pandas as pd
from docx import Document
import time
import comtypes.client

# Load Excel file
excel_file = "./data/example.xlsx"
df = pd.read_excel(excel_file)

# Load the Word template
template_path = "./templates/certificate.docx"
output_dir = "certificates_docx"
pdf_output_dir = "certificates_pdf"

os.makedirs(output_dir, exist_ok=True)
os.makedirs(pdf_output_dir, exist_ok=True)

# Function to convert DOCX to PDF
def convert_to_pdf(docx_path, pdf_path):
    word = comtypes.client.CreateObject("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(os.path.abspath(docx_path))
    doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)  # 17 = PDF
    doc.Close()
    word.Quit()
    print(f"✅ Converted {docx_path} to {pdf_path}")

# Generate certificates
for _, row in df.iterrows():
    first_name = str(row["first_name"])
    last_name = str(row["last_name"])
    date = pd.Timestamp.now().strftime("%B %d, %Y")

    # Load template
    doc = Document(template_path)

    # Replace placeholders
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace("{{FIRST_NAME}}", first_name)
            run.text = run.text.replace("{{LAST_NAME}}", last_name)
            run.text = run.text.replace("{{DATE}}", date)
        # print([run.text for run in paragraph.runs]) #Used to test the output before generating docs

    # Save the DOCX certificate
    docx_file = os.path.join(output_dir, f"Certificate-{first_name}_{last_name}.docx")
    doc.save(docx_file)

    time.sleep(2)

    # Convert to PDF
    pdf_file = os.path.join(pdf_output_dir, f"Certificate-{first_name}_{last_name}.pdf")
    convert_to_pdf(docx_file, pdf_file)

print("✅ Certificates generated and converted to PDF successfully!")
